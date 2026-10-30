"""Mobi Price Finder — controlled text price-list imports into PostgreSQL."""
import html
import logging
import os
import re
import hashlib
import asyncio
import tempfile
from io import BytesIO
from pathlib import Path
from large_csv_importer import import_large_csv, MAX_CSV_BYTES
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import datetime, timedelta
import time
from datetime import timezone
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.error import BadRequest, TelegramError
from telegram.ext import (
    ApplicationBuilder, CallbackQueryHandler, CommandHandler,
    ContextTypes, MessageHandler, filters,
)
from myfin_rate import MYFIN_URL, fetch_bsb_usd_sale
from mts_daily_scheduler import post_init as mts_post_init
import a1_bot_integration
from database import (
    add_supplier, get_supplier, get_supplier_offers, get_offers, search_offers, init_db, list_suppliers,
    save_price_list, save_excel_rows, update_supplier_settings, rename_supplier, delete_supplier,
    get_excel_profiles, save_excel_profiles, export_supplier_offers, iter_supplier_offers,
    get_usd_byn_rate, set_usd_byn_rate, set_supplier_delivery_days,
)
from product_parser import parse_product, section_brand, contextual_name
from excel_importer import (inspect_xlsx, extract_xlsx, MAX_FILE_BYTES,
                            sheet_columns, inspect_mapping, sheet_fingerprint)

logging.basicConfig(format="%(asctime)s - %(levelname)s - %(message)s", level=logging.INFO)
logging.getLogger("httpx").setLevel(logging.WARNING)
logger = logging.getLogger(__name__)


def _user_id(update):
    user = update.effective_user
    return str(user.id) if user else ""


def is_owner(update):
    owner_id = os.getenv("OWNER_TELEGRAM_ID", "").strip()
    return bool(owner_id and owner_id.isdigit() and _user_id(update) == owner_id)


def is_manager(update):
    # One or several IDs, separated by comma, semicolon or whitespace.
    raw_ids = os.getenv("MANAGER_TELEGRAM_IDS", "")
    manager_ids = set(re.findall(r"\d+", raw_ids))
    return bool(_user_id(update) and _user_id(update) in manager_ids)


def allowed(update):
    return is_owner(update) or is_manager(update)


async def authorize(update):
    if allowed(update):
        return True
    await update.effective_message.reply_text(
        "🔒 Доступ закрыт. Отправь /myid и передай Telegram ID владельцу бота "
        "для добавления в MANAGER_TELEGRAM_IDS."
    )
    return False


def _clean_price_name(value):
    # Remove decorative symbols, not letters or meaningful model attributes.
    value = re.sub(r"^[\s•▪▫◾🔹🔸✅☑️📱🎧⌚💻📦💎🔥⭐️✨➖➡️—–-\uFE0F]+", "", value)
    return value.strip(" \t|:—–-•")


def _supplier_default_brand(supplier_id):
    """Optional Apple-only supplier IDs, e.g. APPLE_ONLY_SUPPLIER_IDS=3,7."""
    ids = set(re.findall(r"\d+", os.getenv('APPLE_ONLY_SUPPLIER_IDS', '')))
    return 'Apple' if str(supplier_id) in ids else None


def multipart_keyboard():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("✅ Закончить и проверить", callback_data="parts:finish")],
        [InlineKeyboardButton("❌ Отменить сбор", callback_data="parts:cancel")],
    ])


def multipart_start_keyboard():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("📚 Загрузить прайс частями", callback_data="parts:start")],
    ])


def parse_price_list(text, initial_brand=None, return_brand=False):
    """Parse supplier text using an explicit terminal price, preserving headings.

    A bare final number is accepted only when the prefix looks like a product
    and the price is not confused with an isolated model number or quantity.
    """
    output = []
    pending = None
    active_brand = initial_brand
    number = r"(\d+(?:[.,]\d{1,2})?)"
    currency = r"(?:\$|USD|EUR|RUB|₽|€)?"
    trailing_note = r"(?:\s*\((?:e\s*sim|sim\s*\+\s*esim|dual\s*sim)\))?"
    # Miscellaneous availability notes follow the price in real supplier lists.
    last_note = r"(?:\s*(?:🔥|бронь|завтра|в наличии|резерв|\((?:актив(?:,\s*запак)?|бронь|резерв)\)))*"
    priced = re.compile(rf"^(.+?)\s*[-—–:]\s*[$€₽]?\s*{number}\s*{currency}{trailing_note}{last_note}\s*$", re.I)
    standalone = re.compile(rf"^[$€₽]?\s*{number}\s*{currency}{trailing_note}{last_note}\s*$", re.I)
    bare = re.compile(rf"^(.+?)\s+[$€₽]?\s*{number}\s*{currency}{trailing_note}{last_note}\s*$", re.I)
    section_labels = re.compile(r"^(?:iphone|iphones|apple|samsung|redmi|xiaomi|macbook|ipad|apple watch|dyson)\s*[:：]?\s*[\W_]*$", re.I)
    separators = re.compile(r"^[\W_]+$", re.UNICODE)

    def is_product_name(name, price, explicit=False):
        if not name or not any(ch.isalpha() for ch in name):
            return False
        if re.search(r"(?i)^(?:от\s+\d+\s*шт|доставка|оплата|заказ до|курс|скидка|отдаю с)\b", name):
            return False
        # A trailing generation number is not necessarily a price: "AirPods 4".
        if not explicit and float(price.replace(',', '.')) < 10:
            return False
        if not explicit and len(name.split()) < 2 and not re.search(r'[\u0400-\u04ff]', name):
            return False
        return True

    def clean_line(line):
        # Markdown emphasis in copied supplier messages is not product data.
        return html.unescape(line).replace('**', '').replace('`', '').replace('\u00a0', ' ').strip().strip('\\*_ ')

    for raw in text.splitlines():
        line = clean_line(raw)
        if not line:
            continue
        heading = section_brand(line)
        if heading:
            active_brand, pending = heading, None
            continue
        if section_labels.fullmatch(line):
            key = re.sub(r'[^a-z]', '', line.lower())
            if key in ('iphone', 'iphones', 'apple'):
                active_brand = 'Apple'
            elif key in ('redmi', 'xiaomi', 'samsung'):
                active_brand = key.capitalize()
            pending = None
            continue
        # A product-specific subheading is not a SKU.
        if re.match(r'^\d{1,2}\s+pro\s*/\s*pro\s*max\b', line, re.I):
            continue
        if separators.fullmatch(line) or line.startswith('@'):
            pending = None
            continue
        # Avoid incorrectly choosing between prices for separate regions.
        # Example: "🇭🇰/🇪🇺 1250/1270" needs a separate future workflow.
        if re.search(r"\b\d{2,6}\s*/\s*\d{2,6}\s*(?:🔥|бронь)?\s*$", line):
            pending = None
            continue
        match = priced.fullmatch(line)
        explicit = bool(match)
        if not match:
            match = bare.fullmatch(line)
        if match:
            name = _clean_price_name(match.group(1))
            price = match.group(2)
            if is_product_name(name, price, explicit):
                # Preserve eSIM when the supplier puts the note after the price.
                if re.search(r'\(e\s*sim\)', line, re.I) and not re.search(r'\besim\b', name, re.I):
                    name += ' eSIM'
                if re.search(r'\(актив(?:,\s*запак)?\)', line, re.I):
                    name += ' Активирован'
                # Unlabelled iPhone generations are frequent after non-iPhone
                # categories. Only infer Apple from strongly iPhone-shaped rows;
                # explicit Samsung/Redmi/Xiaomi context takes precedence.
                brand = active_brand
                if brand not in ('Samsung', 'Redmi', 'Xiaomi') and re.match(
                    r'^\d{2}(?:e|\s+(?:pro(?:\s+max)?|plus))?\s+'
                    r'(?:64|128|256|512|1\s*tb|2\s*tb)\b', name, re.I
                ):
                    brand = 'Apple'
                output.append((contextual_name(name, brand), price))
            pending = None
            continue
        match = standalone.fullmatch(line)
        if match and pending:
            name = _clean_price_name(pending)
            if is_product_name(name, match.group(1), True):
                output.append((contextual_name(name, active_brand), match.group(1)))
            pending = None
            continue
        # Do not take bare prices from paragraphs such as "+1%".
        pending = line if len(line) < 250 and not re.search(r'\bпо запросу\b', line, re.I) else None
    return (output, active_brand) if return_brand else output


def auto_draft_keyboard(draft):
    revision = draft['revision']
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("✅ Сохранить общий прайс", callback_data=f"draft:save:{revision}")],
        [InlineKeyboardButton("❌ Отменить сбор", callback_data=f"draft:cancel:{revision}")],
    ])


def append_to_auto_draft(context, supplier_id, incoming):
    """Reparse complete joined text so headings and broken line boundaries survive."""
    draft = context.user_data.get('auto_draft')
    if draft and draft['supplier_id'] != supplier_id:
        raise ValueError('Сначала сохрани или отмени черновик другого поставщика.')
    if draft is None:
        draft = {'supplier_id': supplier_id, 'messages': [], 'revision': 0, 'rows': []}
    if len(draft['messages']) >= 100 or sum(map(len, draft['messages'])) + len(incoming) > 800_000:
        raise ValueError('Достигнут лимит черновика. Сохрани его или отмени.')
    if incoming in draft['messages']:
        return draft, False
    messages = draft['messages'] + [incoming]
    joined = '\n'.join(messages)
    rows = parse_price_list(joined, _supplier_default_brand(supplier_id))
    if len(rows) > 25_000:
        raise ValueError('Слишком много товарных строк в черновике.')
    draft.update(messages=messages, rows=rows, revision=draft['revision'] + 1)
    context.user_data['auto_draft'] = draft
    context.user_data.pop('pending', None)
    context.user_data.pop('undecided_text', None)
    return draft, True


async def auto_draft_status(message, draft, context):
    """Keep one live Telegram preview instead of one message per pasted part."""
    body = (
        f"📥 Черновик прайса: {len(draft['messages'])} сообщ.\n"
        f"Распознано строк с ценой: {len(draft['rows'])}.\n"
        "Можешь отправить следующие сообщения. Когда всё пришло, нажми «Сохранить общий прайс». "
        "До нажатия данные в базе не изменятся."
    )
    status_chat_id = draft.get('status_chat_id')
    status_message_id = draft.get('status_message_id')
    if status_chat_id and status_message_id:
        try:
            await context.bot.edit_message_text(
                chat_id=status_chat_id, message_id=status_message_id,
                text=body, reply_markup=auto_draft_keyboard(draft),
            )
            return
        except BadRequest as exc:
            # A deleted/inaccessible old preview is recoverable: create a new one.
            if 'message is not modified' in str(exc).lower():
                return
            logger.warning('Cannot edit draft preview, sending new one: %s', exc)
        except TelegramError as exc:
            logger.warning('Cannot edit draft preview, sending new one: %s', exc)
    sent = await message.reply_text(body, reply_markup=auto_draft_keyboard(draft))
    draft['status_chat_id'] = sent.chat_id
    draft['status_message_id'] = sent.message_id


def mode_keyboard():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("💱 Курс USD/BYN", callback_data="fx:menu")],
        [InlineKeyboardButton("🔄 Авто", callback_data="mode:auto"),
         InlineKeyboardButton("🔎 Поиск", callback_data="mode:search")],
        [InlineKeyboardButton("📥 Загрузка прайса", callback_data="mode:import")],
        [InlineKeyboardButton("👥 Поставщики", callback_data="nav:suppliers"),
         InlineKeyboardButton("⚙️ Настройки", callback_data="nav:settings")],
    ])


def decision_keyboard():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🔎 Искать", callback_data="decide:search"),
         InlineKeyboardButton("📥 Это прайс", callback_data="decide:import")],
    ])


async def mode(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await authorize(update):
        return
    current = context.user_data.get('mode', 'auto')
    labels = {'auto': 'Авто', 'search': 'Поиск', 'import': 'Загрузка прайса'}
    await update.effective_message.reply_text(
        f"Режим: {labels.get(current, 'Авто')}. Выбери нужный:",
        reply_markup=mode_keyboard(),
    )


async def menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await mode(update, context)


def _price_import_time(value, short=False):
    if value is None:
        return 'нет данных'
    try:
        tz_name = os.getenv('PRICE_TIMEZONE', 'Europe/Moscow')
        try:
            tz = ZoneInfo(tz_name)
        except ZoneInfoNotFoundError:
            tz = timezone.utc
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        return value.astimezone(tz).strftime('%d.%m' if short else '%d.%m.%Y в %H:%M')
    except (AttributeError, ValueError, TypeError):
        return 'нет данных'


def _delivery_label(days):
    if days is None:
        return 'Не указан'
    if days == 0:
        return 'Сегодня (0 рабочих дней)'
    return f'{days} раб. дн.'


def _delivery_date(days):
    if days is None:
        return None
    try:
        tz = ZoneInfo(os.getenv('PRICE_TIMEZONE', 'Europe/Moscow'))
    except ZoneInfoNotFoundError:
        tz = timezone.utc
    day = datetime.now(tz).date()
    count = 0
    while count < days:
        day += timedelta(days=1)
        if day.weekday() < 5:
            count += 1
    return day.strftime('%d.%m.%Y')


def _supplier_title(supplier):
    return (f"{supplier['name']}\n"
            f"🕒 Последний прайс: {_price_import_time(supplier.get('last_price_import_at'))}\n"
            f"🚚 Доставка: {_delivery_label(supplier.get('delivery_workdays'))}")


def keyboard_suppliers():
    buttons = [[InlineKeyboardButton(
                    f"{s['name']} · {_price_import_time(s.get('last_price_import_at'), short=True)}",
                    callback_data=f"pick:{s['id']}")]
               for s in list_suppliers()[:75]]
    buttons.append([InlineKeyboardButton("➕ Новый поставщик", callback_data="supplier:new")])
    return InlineKeyboardMarkup(buttons)


def supplier_manage_keyboard(supplier, can_delete=True):
    sid = supplier['id']
    buttons = [[InlineKeyboardButton("✏️ Переименовать", callback_data=f"supplier:rename:{sid}")]]
    if can_delete:
        buttons.append([InlineKeyboardButton("🗑 Удалить", callback_data=f"supplier:delete:{sid}")])
    buttons += [
        [InlineKeyboardButton("📥 Выгрузить прайс (Excel)", callback_data=f"supplier:export:{sid}")],
        [InlineKeyboardButton("⚙️ Валюта, НДС и доставка", callback_data=f"supplier:settings:{sid}")],
        [InlineKeyboardButton("👥 Все поставщики", callback_data="nav:suppliers")],
    ]
    return InlineKeyboardMarkup(buttons)


def supplier_delete_keyboard(sid):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("🗑 Да, удалить с товарами", callback_data=f"supplier:confirm_delete:{sid}")],
        [InlineKeyboardButton("↩️ Отмена", callback_data=f"supplier:cancel_delete:{sid}")],
    ])


def settings_keyboard(supplier):
    sid = supplier['id']
    buttons = [
        [InlineKeyboardButton(f"Валюта: {supplier['currency']}", callback_data=f"settings:currency:{sid}")],
        [InlineKeyboardButton(f"НДС: {supplier['tax_status']}", callback_data=f"settings:vat:{sid}")],
        [InlineKeyboardButton(f"🚚 Доставка: {_delivery_label(supplier.get('delivery_workdays'))}", callback_data=f"settings:delivery:{sid}")],
    ]
    return InlineKeyboardMarkup(buttons)


def confirmation_keyboard():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("✅ Сохранить прайс", callback_data="import:save")],
        [InlineKeyboardButton("❌ Отменить", callback_data="import:cancel")],
    ])


async def myid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.effective_message.reply_text(f"Твой Telegram ID: {update.effective_user.id}")


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await authorize(update):
        return
    await update.effective_message.reply_text(
        "👋 Mobi Price Finder\n\n/suppliers — выбрать/добавить поставщика\n"
        "/settings — валюта и НДС выбранного поставщика\n"
        "/prices — сохранённые цены выбранного поставщика\n"
        "/find iPhone 17 Pro 256GB Black — сравнить поставщиков\n"
        "/menu — меню и режим работы\n"
        "Просто напиши название товара — бот выполнит поиск.\n"
        "Прайсы с ценами бот предложит сохранить.\n"
        "Для длинного текстового прайса выбери 📚 Прайс частями.",
        reply_markup=mode_keyboard(),
    )


async def suppliers(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await authorize(update):
        return
    await update.effective_message.reply_text("Выбери поставщика (рядом дата последнего прайса):", reply_markup=keyboard_suppliers())


async def addsupplier(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await authorize(update):
        return
    context.user_data['await_name'] = True
    await update.effective_message.reply_text("Введи название нового поставщика одним сообщением.")


async def settings(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await authorize(update):
        return
    sid = context.user_data.get('supplier_id')
    s = get_supplier(sid) if sid else None
    if not s:
        await update.effective_message.reply_text("Сначала выбери поставщика через /suppliers")
        return
    await update.effective_message.reply_text(
        f"⚙️ {s['name']}\nВалюта: {s['currency']}\nНДС: {s['tax_status']}\nДоставка: {_delivery_label(s.get('delivery_workdays'))}",
        reply_markup=settings_keyboard(s),
    )


async def prices(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await authorize(update):
        return
    sid = context.user_data.get('supplier_id')
    supplier = get_supplier(sid) if sid else None
    if supplier is None:
        await update.effective_message.reply_text(
            'Сначала выбери поставщика через /suppliers.'
        )
        return

    offers = get_supplier_offers(supplier['id'], limit=20)
    if not offers:
        await update.effective_message.reply_text(
            f"У поставщика «{supplier['name']}» пока нет сохранённых предложений."
        )
        return

    lines = [f"📦 {supplier['name']} — последние {len(offers)} предложений:"]
    for i, offer in enumerate(offers, 1):
        price = (f"{offer['price']} {offer['currency'] or 'валюта неизвестна'}"
                 if offer['price'] is not None else 'По запросу')
        fields = []
        if offer['region']:
            fields.append(f"регион: {offer['region']}")
        if offer['sim']:
            fields.append(f"SIM: {offer['sim']}")
        if offer['tax_status'] == 'with_vat':
            fields.append('с НДС')
        elif offer['tax_status'] == 'without_vat':
            fields.append('без НДС')
        freshness = 'актуально' if offer['is_fresh'] else 'срок актуальности истёк'
        date = offer['updated_at'].strftime('%d.%m.%Y %H:%M UTC')
        details = ' | '.join(fields)
        lines.append(
            f"\n{i}. {offer['original_name']}\n"
            f"💰 {price}\n"
            + (f"{details}\n" if details else '')
            + f"🕒 {date} · {freshness}"
        )

    # Telegram limits messages to 4096 characters.
    text = ''
    for line in lines:
        if len(text) + len(line) + 1 > 3500:
            if text:
                await update.effective_message.reply_text(text)
            text = ''
        text += ('\n' if text else '') + line
    if text:
        await update.effective_message.reply_text(text)


async def show_pending(message, context):
    pending = context.user_data.get('pending')
    if not pending:
        return
    s = get_supplier(context.user_data.get('supplier_id')) if context.user_data.get('supplier_id') else None
    if not s:
        await message.reply_text("Выбери поставщика для прайса:", reply_markup=keyboard_suppliers())
        return
    if s['currency'] not in ('USD', 'BYN', 'RUB', 'EUR'):
        await message.reply_text("⚠️ Сначала выбери валюту в /settings.")
        return
    if s['tax_status'] == 'unknown':
        await message.reply_text(
            f"⚠️ У поставщика {s['name']} не указан НДС. Открой /settings и установи его "
            "перед сохранением."
        )
        return
    parsed = [(parse_product(name), price) for name, price in pending]
    supported = parsed
    previews = [f"• {p['original']} — {v}" for p, v in supported[:3]]
    await message.reply_text(
        f"📋 Прайс: {len(parsed)} строк\nПоставщик: {s['name']}\n"
        f"Валюта: {s['currency']} | НДС: {s['tax_status']}\n"
        f"Строк для импорта: {len(supported)} (точные дубли названий свернём)\n"
        f"Нераспознанные модели сохраняются по исходному названию\n\n"
        + "\n".join(previews)
        + "\n\nСохранить?",
        reply_markup=confirmation_keyboard(),
    )



EXCEL_FIELDS = ('name', 'price', 'brand', 'color', 'region', 'category')
EXCEL_FIELD_NAMES = {'name': 'Название товара', 'price': 'Цена', 'brand': 'Бренд',
                     'color': 'Цвет', 'region': 'Регион', 'category': 'Категория'}


def excel_menu(stage):
    buttons = []
    for i, item in enumerate(stage['sheets']):
        mark = '✅' if item.get('enabled') else '⏭' if item.get('skipped') else '⚠️'
        buttons.append([InlineKeyboardButton(f'{mark} {item.get("file_name", "")[:16]} · {item["sheet"][:20]}', callback_data=f'xm:sheet:{i}')])
    buttons.append([InlineKeyboardButton('✅ Предпросмотр и сохранение', callback_data='xm:review')])
    buttons.append([InlineKeyboardButton('❌ Отмена', callback_data='excel:cancel')])
    return InlineKeyboardMarkup(buttons)


def excel_sheet_menu(stage, index):
    item = stage['sheets'][index]
    mapping = item.get('mapping') or {}
    def letter(field):
        idx = mapping.get(field)
        return '-' if idx is None else __import__('openpyxl').utils.get_column_letter(idx+1)
    buttons = [[InlineKeyboardButton(f'Строка заголовков: {item.get("header_row", 0)}', callback_data=f'xm:head:{index}')]]
    for field in EXCEL_FIELDS:
        buttons.append([InlineKeyboardButton(f'{EXCEL_FIELD_NAMES[field]}: {letter(field)}', callback_data=f'xm:field:{index}:{field}')])
    buttons.extend([[InlineKeyboardButton('✅ Использовать лист', callback_data=f'xm:use:{index}')],
                    [InlineKeyboardButton('⏭ Пропустить лист', callback_data=f'xm:skip:{index}')],
                    [InlineKeyboardButton('↩️ К списку листов', callback_data='xm:list')]])
    return InlineKeyboardMarkup(buttons)


def staged_excel_ready(stage):
    sheets = stage['sheets']
    if not all(item.get('enabled') or item.get('skipped') for item in sheets):
        return False
    if not stage.get('files'):
        return any(item.get('enabled') for item in sheets)
    # Each attached file must contribute at least one enabled sheet.
    return all(any(sheet.get('enabled') for sheet in attached['sheets'])
               for attached in stage['files'])


async def excel_review(message, context):
    stage = context.user_data.get('excel_pending')
    if not stage or stage['expires'] < time.time():
        await message.reply_text('Настройка истекла. Пришли Excel ещё раз.')
        return
    lines = ['📊 Excel — проверка перед импортом', 'Поставщик: '+stage['supplier_name'], f'Файлов: {len(stage.get("files", [1]))}']
    for item in stage['sheets']:
        if item.get('enabled'):
            lines.append(f'✅ {item.get("file_name", "Excel")} / {item["sheet"]}: {item["valid"]} строк')
            lines.extend('• '+v for v in item.get('samples', [])[:2])
            for parsed in item.get('parsed_samples', [])[:2]:
                display = lambda v: str(v) if v is not None else 'Не указан'
                lines.append('   ↳ Модель: '+display(parsed.get('model'))
                             +' | Память: '+(str(parsed['storage'])+' GB' if parsed.get('storage') is not None else 'Не указан'))
                lines.append('   ↳ Цвет: '+display(parsed.get('color'))
                             +' | SIM: '+display(parsed.get('sim'))
                             +' | Регион: '+display(parsed.get('region')))
                if parsed.get('warnings'):
                    lines.append('   ⚠️ '+'; '.join(parsed['warnings'])[:130])
            if item.get('uncertain'):
                if item.get('uncertain_sampled'):
                    lines.append(f'⚠️ Неоднозначные характеристики в выборке первых 100 товаров: {item["uncertain"]}. Исходные названия сохраняются.')
                else:
                    lines.append(f'⚠️ Неоднозначные характеристики: {item["uncertain"]} строк(и). Исходные названия сохраняются.')
        elif item.get('skipped'):
            lines.append('⏭ '+item.get('file_name', 'Excel')+' / '+item['sheet']+' — пропущен')
        else:
            lines.append('⚠️ '+item.get('file_name', 'Excel')+' / '+item['sheet']+' — требует настройки')
    lines.append('Никакие товары ещё не сохранены. При подтверждении прежний прайс выбранного поставщика будет заменён.')
    keyboard = [[InlineKeyboardButton('🛠 Настроить листы', callback_data='xm:list')]]
    if staged_excel_ready(stage):
        keyboard.insert(0, [InlineKeyboardButton('✅ Сохранить прайс', callback_data='excel:save')])
    keyboard.append([InlineKeyboardButton('❌ Отменить', callback_data='excel:cancel')])
    await message.reply_text('\n'.join(lines)[:4000], reply_markup=InlineKeyboardMarkup(keyboard))


async def excel_manual_callback(query, context, action):
    stage = context.user_data.get('excel_pending')
    if not stage or stage['expires'] < time.time():
        await query.message.reply_text('Настройка Excel истекла. Загрузи файл ещё раз.')
        return
    parts = action.split(':')
    op = parts[1]
    if op == 'review':
        await excel_review(query.message, context)
        return
    if op == 'list':
        await query.message.reply_text('Выбери лист для настройки. ✅ — включён, ⏭ — пропущен, ⚠️ — требует настройки.', reply_markup=excel_menu(stage))
        return
    try:
        idx = int(parts[2]); item = stage['sheets'][idx]
    except (IndexError, ValueError):
        return
    if op == 'sheet':
        await query.message.reply_text(f'Лист «{item["sheet"]}». Выбери заголовки и колонки, затем нажми «Использовать лист».', reply_markup=excel_sheet_menu(stage, idx))
    elif op == 'head':
        choices = [[InlineKeyboardButton('Без заголовков (0)', callback_data=f'xm:sethead:{idx}:0')]]
        choices += [[InlineKeyboardButton(f'Строка {n}', callback_data=f'xm:sethead:{idx}:{n}')] for n in range(1, 21)]
        await query.message.reply_text('Выбери номер строки с заголовками. 0 — данные начинаются с первой строки.', reply_markup=InlineKeyboardMarkup(choices))
    elif op == 'sethead':
        item['header_row'] = int(parts[3]); item['enabled'] = False; item['skipped'] = False
        await query.message.reply_text('Строка установлена.', reply_markup=excel_sheet_menu(stage, idx))
    elif op == 'field':
        field = parts[3]
        if field not in EXCEL_FIELDS:
            return
        columns = sheet_columns(item.get('data', stage['data']), item['sheet'])
        choices = []
        if field not in ('name', 'price'):
            choices.append([InlineKeyboardButton('Не использовать', callback_data=f'xm:setcol:{idx}:{field}:x')])
        for col in columns:
            descr = ' / '.join(col['examples'])[:37] or '(пусто в начале)'
            choices.append([InlineKeyboardButton(f'{col["letter"]} — {descr}', callback_data=f'xm:setcol:{idx}:{field}:{col["index"]}')])
        await query.message.reply_text('Выбери столбец: '+EXCEL_FIELD_NAMES[field], reply_markup=InlineKeyboardMarkup(choices))
    elif op == 'setcol':
        field, chosen = parts[3], parts[4]
        if field not in EXCEL_FIELDS:
            return
        item.setdefault('mapping', {k: None for k in EXCEL_FIELDS})[field] = None if chosen == 'x' else int(chosen)
        item['enabled'] = False; item['skipped'] = False
        await query.message.reply_text('Столбец выбран.', reply_markup=excel_sheet_menu(stage, idx))
    elif op == 'skip':
        item['enabled'] = False; item['skipped'] = True
        await query.message.reply_text('Лист пропущен.', reply_markup=excel_menu(stage))
    elif op == 'use':
        mapping = item.get('mapping') or {}
        if mapping.get('name') is None or mapping.get('price') is None or mapping['name'] == mapping['price']:
            await query.message.reply_text('❌ Укажи разные столбцы товара и цены.')
            return
        info = inspect_mapping(item.get('data', stage['data']), item['sheet'], mapping, item.get('header_row', 0))
        if not info['valid'] or info['truncated']:
            await query.message.reply_text('❌ Нет строк с распознанными названиями и положительными ценами либо лист слишком велик. Исправь столбцы.')
            return
        item.update(info); item['enabled'] = True; item['skipped'] = False
        await query.message.reply_text(f'✅ {item["sheet"]}: найдено {info["valid"]} предложений.', reply_markup=excel_menu(stage))


def _audit_cell(value):
    """Prevent an imported supplier name from becoming an Excel formula."""
    if value is None:
        return 'Не указан'
    if isinstance(value, str):
        value = value.strip()
        return "'" + value if value.startswith(('=', '+', '-', '@')) else (value or 'Не указан')
    return value


# XLSX limit: 1,048,576 rows per sheet, including header.
_XLSX_DATA_ROWS_PER_SHEET = 1_048_575
_supplier_export_lock = asyncio.Lock()


def _build_supplier_export(supplier, offers, path):
    """Write an Excel audit directly to disk with bounded RAM usage.

    The caller supplies a streaming iterator of database records. Both sheets
    split at Excel's row limit; original column order and audit rules remain.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)
    headers = [
        'ID записи', 'Исходное название (БД)', 'Модель в поиске', 'Бренд',
        'Память, GB', 'RAM, GB', 'Цвет', 'Регион', 'SIM', 'Состояние',
        'Цена', 'Валюта', 'НДС', 'Обновлён', 'Актуален',
        'Модель сохранена в БД', 'Цвет сохранён в БД', 'Регион сохранён в БД',
        'SIM сохранён в БД', 'Предупреждения БД',
    ]
    check_headers = ['ID записи', 'Исходное название', 'Причина проверки', 'Цена', 'Валюта']
    all_widths = [13, 70, 29, 17, 15, 14, 22, 21, 18, 20, 16, 13, 16, 22, 12, 29, 22, 21, 18, 50]
    check_widths = [13, 75, 65, 16, 14]
    header_fill = PatternFill('solid', fgColor='17365D')
    header_font = Font(color='FFFFFF', bold=True)
    header_alignment = Alignment(vertical='center', wrap_text=True)
    price_format = '#,##0.00'

    def create_sheet(base, number, headings, widths, price_column):
        title = base if number == 1 else f'{base} {number}'
        sheet = workbook.create_sheet(title)
        sheet.freeze_panes = 'C2' if base == 'Все товары' else 'B2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(headings))}{_XLSX_DATA_ROWS_PER_SHEET + 1}'
        for idx, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        cells = []
        for heading in headings:
            cell = WriteOnlyCell(sheet, value=heading)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cells.append(cell)
        sheet.append(cells)
        return sheet

    all_number = check_number = 1
    all_sheet = create_sheet('Все товары', all_number, headers, all_widths, 11)
    check_sheet = create_sheet('Проверка', check_number, check_headers, check_widths, 4)
    total = review_count = 0
    vat_labels = {'with_vat': 'С НДС', 'without_vat': 'Без НДС', 'unknown': 'Не указан'}
    condition_labels = {'activated': 'Активирован', 'sealed': 'Запечатан',
                        'refurbished': 'Восстановлен', 'used': 'Б/у'}

    def format_row(sheet, values, price_col, date_col=None):
        # Create styled cells only for columns that need a number format.
        values = list(values)
        for index, number_format in ((price_col - 1, price_format),
                                      (date_col - 1, 'DD.MM.YYYY HH:MM') if date_col else (-1, '')):
            if index >= 0 and values[index] is not None:
                cell = WriteOnlyCell(sheet, value=values[index])
                cell.number_format = number_format
                values[index] = cell
        sheet.append(values)

    for offer in offers:
        if total and total % _XLSX_DATA_ROWS_PER_SHEET == 0:
            all_number += 1
            all_sheet = create_sheet('Все товары', all_number, headers, all_widths, 11)
        parsed = parse_product(offer['original_name'])
        def effective(key):
            value = parsed.get(key)
            return value if value is not None and value != [] and value != '' else offer.get(key)
        conditions = parsed.get('condition') or offer.get('condition_tags') or []
        if isinstance(conditions, str):
            conditions = [conditions]
        status = ', '.join(condition_labels.get(str(c), str(c)) for c in conditions) or 'Не указано'
        date = offer.get('updated_at')
        if date and date.tzinfo:
            date = date.astimezone(ZoneInfo(os.getenv('PRICE_TIMEZONE', 'Europe/Moscow'))).replace(tzinfo=None)
        price = float(offer['price']) if offer.get('price') is not None else None
        row = [
            offer['id'], _audit_cell(offer['original_name']), _audit_cell(effective('model')),
            _audit_cell(effective('brand')), effective('storage_gb'), effective('ram_gb'),
            _audit_cell(effective('color')), _audit_cell(effective('region')),
            _audit_cell(effective('sim')), status, price, offer.get('currency'),
            vat_labels.get(offer.get('tax_status'), 'Не указан'), date,
            'Да' if offer.get('is_fresh') else 'Нет',
            _audit_cell(offer.get('model')), _audit_cell(offer.get('color')),
            _audit_cell(offer.get('region')), _audit_cell(offer.get('sim')),
            '; '.join(str(x) for x in (offer.get('warnings') or [])),
        ]
        format_row(all_sheet, row, 11, 14)
        total += 1
        if total % 100000 == 0:
            logger.info('Excel export: %s rows written for supplier %s', total, supplier['id'])
        reasons = []
        if not effective('model'):
            reasons.append('Модель не распознана')
        if price is None or price <= 0:
            reasons.append('Некорректная или отсутствующая цена')
        if price is not None and price < 50:
            reasons.append('Цена ниже 50 — проверь, если это не аксессуар')
        if not offer.get('is_fresh'):
            reasons.append('Предложение устарело')
        if parsed.get('model') and offer.get('model') and parsed['model'] != offer['model']:
            reasons.append('Модель из текста отличается от сохранённой')
        reasons += [str(w) for w in (offer.get('warnings') or [])]
        if reasons:
            if review_count and review_count % _XLSX_DATA_ROWS_PER_SHEET == 0:
                check_number += 1
                check_sheet = create_sheet('Проверка', check_number, check_headers, check_widths, 4)
            format_row(check_sheet, [offer['id'], _audit_cell(offer['original_name']),
                                     '; '.join(reasons), price, offer.get('currency')], 4)
            review_count += 1

    if total:
        workbook.save(path)
    workbook.close()
    return total, review_count


async def callbacks(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    if not allowed(update):
        await query.answer("Нет доступа", show_alert=True)
        return
    await query.answer()
    action = query.data or ''
    if action.startswith('sr:'):
        await search_results_callback(query, context, action)
        return
    if action.startswith('draft:'):
        op, revision_raw = action.split(':')[1:]
        draft = context.user_data.get('auto_draft')
        if not draft or not revision_raw.isdigit() or int(revision_raw) != draft['revision']:
            await query.message.reply_text('Кнопка устарела: используй кнопки в последнем сообщении о черновике.')
            return
        if op == 'cancel':
            context.user_data.pop('auto_draft', None)
            try:
                await context.bot.edit_message_reply_markup(
                    chat_id=draft.get('status_chat_id'), message_id=draft.get('status_message_id'),
                    reply_markup=None,
                )
            except TelegramError:
                pass
            await query.message.reply_text('❌ Черновик отменён. Сохранённый прайс не изменён.')
        elif op == 'save':
            sid = context.user_data.get('supplier_id')
            if sid != draft['supplier_id'] or not get_supplier(sid):
                await query.message.reply_text('Поставщик изменён. Выбери прежнего поставщика.')
                return
            if not draft['rows']:
                await query.message.reply_text('Нет распознанных товаров. Прайс не заменён.')
                return
            try:
                parsed = [(parse_product(name), price) for name, price in draft['rows']]
                summary = save_price_list(sid, parsed)
            except Exception:
                logger.exception('Failed to save auto draft')
                await query.message.reply_text('❌ Ошибка сохранения. Старый прайс должен остаться без изменений. Проверь Railway Logs.')
                return
            context.user_data.pop('auto_draft', None)
            try:
                await context.bot.edit_message_reply_markup(
                    chat_id=draft.get('status_chat_id'), message_id=draft.get('status_message_id'),
                    reply_markup=None,
                )
            except TelegramError:
                pass
            await query.message.reply_text(
                f"✅ Общий прайс сохранён!\nСообщений: {len(draft['messages'])}\n"
                f"Конфигураций: {summary['saved']}\nДубли: {summary['duplicates']}\n"
                f"Пропущено: {summary['skipped']}\nСтарый прайс поставщика заменён."
            )
        return
    if action == 'parts:start':
        sid = context.user_data.get('supplier_id')
        supplier = get_supplier(sid) if sid else None
        if not supplier:
            await query.message.reply_text('Сначала выбери поставщика через /suppliers.')
            return
        if context.user_data.get('parts'):
            await query.message.reply_text('Сбор прайса уже идёт. Отправляй следующие части или заверши его.', reply_markup=multipart_keyboard())
            return
        context.user_data['parts'] = {'supplier_id': sid, 'rows': [], 'messages': 0, 'started': time.time(), 'brand': _supplier_default_brand(sid)}
        context.user_data.pop('pending', None)
        context.user_data.pop('undecided_text', None)
        await query.message.reply_text(
            f'📚 Сбор текстового прайса для «{supplier["name"]}».\n'
            'Отправляй части отдельными сообщениями, включая эмодзи. '
            'Ничего не сохраняется до подтверждения.\n'
            'После последней части нажми «Закончить и проверить».',
            reply_markup=multipart_keyboard())
    elif action == 'parts:cancel':
        context.user_data.pop('parts', None)
        await query.message.reply_text('❌ Сбор отменён. Текущий прайс поставщика не изменился.')
    elif action == 'parts:finish':
        batch = context.user_data.get('parts')
        if not batch or not batch['rows']:
            await query.message.reply_text('Нет собранных товаров. Пришли части прайса или отменяй сбор.')
            return
        if context.user_data.get('supplier_id') != batch['supplier_id']:
            await query.message.reply_text('Выбран другой поставщик. Верни прежнего или отмени сбор.')
            return
        context.user_data['pending'] = list(batch['rows'])
        context.user_data.pop('parts', None)
        await show_pending(query.message, context)
    elif action.startswith('mode:'):
        mode_name = action.split(':', 1)[1]
        if mode_name not in ('auto', 'search', 'import'):
            return
        context.user_data['mode'] = mode_name
        labels = {'auto': 'Авто', 'search': 'Поиск', 'import': 'Загрузка прайса'}
        await query.message.reply_text(f"✅ Режим: {labels[mode_name]}")
    elif action == 'fx:menu':
        rate = get_usd_byn_rate()
        buttons = [
            [InlineKeyboardButton('✏️ Установить курс вручную', callback_data='fx:edit')],
            [InlineKeyboardButton('🏦 Получить курс БСБ Банка (Myfin)', callback_data='fx:myfin')],
        ]
        await query.message.reply_text(
            f'💱 Текущий курс: 1 USD = {rate} BYN\n'
            'Источник для загрузки: Myfin, БСБ Банк, продажа USD.\n'
            'Курс меняется только по нажатию кнопки; исходные цены товаров не меняются.',
            reply_markup=InlineKeyboardMarkup(buttons))
    elif action == 'fx:edit':
        context.user_data['await_fx_rate'] = True
        await query.message.reply_text('Введи новый курс: сколько BYN за 1 USD, например 3,05. /cancelrate — отмена.')
    elif action == 'fx:myfin':
        # Fetch the website off the event loop: do not stall the Telegram bot.
        await query.message.reply_text('🏦 Получаю курс продажи USD БСБ Банка с Myfin...')
        try:
            rate = await asyncio.to_thread(fetch_bsb_usd_sale)
            # Only write after the bank has been identified and validated.
            saved = await asyncio.to_thread(set_usd_byn_rate, rate)
        except Exception:
            logger.exception('Failed to update USD/BYN from Myfin')
            await query.message.reply_text(
                '⚠️ Не удалось получить и проверить курс БСБ Банка. '
                'Прежний курс сохранён. Попробуй позднее или установи вручную.')
            return
        await query.message.reply_text(
            f'✅ Курс обновлён: 1 USD = {saved} BYN\n'
            '🏦 БСБ Банк — продажа USD\n'
            f'Источник: {MYFIN_URL}')
    elif action == 'nav:suppliers':
        await query.message.reply_text('Выбери поставщика:', reply_markup=keyboard_suppliers())
    elif action == 'nav:settings':
        sid = context.user_data.get('supplier_id')
        supplier = get_supplier(sid) if sid else None
        if supplier:
            await query.message.reply_text(
                f"⚙️ {supplier['name']}", reply_markup=settings_keyboard(supplier)
            )
        else:
            await query.message.reply_text('Сначала выбери поставщика:', reply_markup=keyboard_suppliers())
    elif action.startswith('decide:'):
        raw = context.user_data.pop('undecided_text', None)
        if not raw:
            await query.message.reply_text('Сообщение для обработки уже отсутствует. Пришли его ещё раз.')
            return
        decision = action.split(':', 1)[1]
        if decision == 'search':
            await run_search(query.message, raw, context)
        elif decision == 'import':
            rows = parse_price_list(raw, _supplier_default_brand(context.user_data.get('supplier_id')))
            if not rows:
                await query.message.reply_text('В сообщении нет распознанных пар «товар — цена».')
            else:
                sid = context.user_data.get('supplier_id')
                if not sid or not get_supplier(sid):
                    await query.message.reply_text('Сначала выбери поставщика через /suppliers.')
                    return
                try:
                    draft, _ = append_to_auto_draft(context, sid, raw)
                except ValueError as exc:
                    await query.message.reply_text(str(exc))
                    return
                await auto_draft_status(query.message, draft, context)
    elif action == 'supplier:new':
        context.user_data.pop('await_rename_id', None)
        context.user_data['await_name'] = True
        await query.message.reply_text("Введи название нового поставщика одним сообщением.")
    elif action.startswith('pick:'):
        sid = int(action.split(':')[1])
        if context.user_data.get('auto_draft') and context.user_data['auto_draft']['supplier_id'] != sid:
            await query.message.reply_text('Сначала сохрани или отмени черновик текущего поставщика.', reply_markup=auto_draft_keyboard(context.user_data['auto_draft']))
            return
        if context.user_data.get('parts') and context.user_data['parts']['supplier_id'] != sid:
            await query.message.reply_text('Сначала заверши или отмени сбор прайса для текущего поставщика.', reply_markup=multipart_keyboard())
            return
        s = get_supplier(sid)
        if not s:
            await query.message.reply_text("Поставщик не найден")
            return
        context.user_data['supplier_id'] = sid
        await query.message.reply_text(f"✅ Выбран поставщик:\n{_supplier_title(s)}", reply_markup=supplier_manage_keyboard(s, allowed(update)))
        await show_pending(query.message, context)
    elif action.startswith('supplier:'):
        parts = action.split(':')
        if len(parts) != 3 or not parts[2].isdigit():
            return
        operation, sid = parts[1], int(parts[2])
        s = get_supplier(sid)
        if not s:
            await query.message.reply_text('Поставщик уже удалён.', reply_markup=keyboard_suppliers())
            return
        if operation == 'export':
            if _supplier_export_lock.locked():
                await query.message.reply_text('⏳ Другая выгрузка Excel уже выполняется. Повтори позже.')
                return
            async with _supplier_export_lock:
                await query.message.reply_text('📦 Формирую Excel. Для большого прайса это может занять несколько минут.')
                try:
                    with tempfile.TemporaryDirectory(prefix='pricefinder_export_') as folder:
                        file_name = f'pricefinder_supplier_{sid}_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
                        path = os.path.join(folder, file_name)
                        total, review_count = await asyncio.to_thread(
                            _build_supplier_export, s, iter_supplier_offers(sid), path)
                        if not total:
                            await query.message.reply_text('У этого поставщика пока нет сохранённых товаров.')
                            return
                        logger.info('Excel export: ready supplier=%s rows=%s review=%s size=%s',
                                    sid, total, review_count, os.path.getsize(path))
                        with open(path, 'rb') as document:
                            await query.message.reply_document(
                                document=document, filename=file_name,
                                read_timeout=600, write_timeout=600,
                                caption=(f'📥 Выгрузка из PostgreSQL: {s["name"]}\n'
                                         f'Товаров: {total}\n'
                                         f'Позиций для проверки: {review_count}\n'
                                         'Показаны все сохранённые товары, включая устаревшие.'),
                            )
                except Exception:
                    logger.exception('Failed to export supplier %s', sid)
                    await query.message.reply_text('Не удалось выгрузить прайс. Проверь логи Railway.')
        elif operation == 'rename':
            context.user_data.pop('await_name', None)
            context.user_data['await_rename_id'] = sid
            await query.message.reply_text(
                f"✏️ Новое название для «{s['name']}»? Отправь одним сообщением. "
                "Чтобы отменить, нажми кнопку.",
                reply_markup=InlineKeyboardMarkup([[
                    InlineKeyboardButton('❌ Отмена', callback_data=f'supplier:cancel_rename:{sid}')
                ]])
            )
        elif operation == 'cancel_rename':
            if context.user_data.get('await_rename_id') == sid:
                context.user_data.pop('await_rename_id', None)
            await query.message.reply_text('Изменение названия отменено.', reply_markup=supplier_manage_keyboard(s, allowed(update)))
        elif operation == 'delete':
            await query.message.reply_text(
                f"⚠️ Удалить «{s['name']}» и ВСЕ его сохранённые товары? "
                "Это действие нельзя отменить.", reply_markup=supplier_delete_keyboard(sid)
            )
        elif operation == 'cancel_delete':
            await query.message.reply_text('Удаление отменено.', reply_markup=supplier_manage_keyboard(s, allowed(update)))
        elif operation == 'confirm_delete':
            deleted = delete_supplier(sid)
            if deleted:
                if context.user_data.get('supplier_id') == sid:
                    context.user_data.pop('supplier_id', None)
                    context.user_data.pop('pending', None)
                    context.user_data.pop('excel_pending', None)
                if context.user_data.get('await_rename_id') == sid:
                    context.user_data.pop('await_rename_id', None)
                await query.message.reply_text(
                    f"✅ Поставщик «{s['name']}» и его товары удалены.",
                    reply_markup=keyboard_suppliers()
                )
            else:
                await query.message.reply_text('Поставщик уже удалён.', reply_markup=keyboard_suppliers())
        elif operation == 'settings':
            context.user_data['supplier_id'] = sid
            await query.message.reply_text(f"⚙️ {s['name']}", reply_markup=settings_keyboard(s))
    elif action.startswith('settings:delivery:'):
        sid = int(action.rsplit(':', 1)[1])
        supplier = get_supplier(sid)
        if not supplier:
            await query.message.reply_text('Поставщик не найден.')
            return
        context.user_data['await_delivery_id'] = sid
        await query.message.reply_text(
            f"🚚 {supplier['name']}\nТекущий срок: {_delivery_label(supplier.get('delivery_workdays'))}\n"
            'Отправь количество РАБОЧИХ дней (0 — сегодня).\n'
            '«нет» — неизвестный срок, «отмена» — не менять.')
    elif action.startswith('settings:currency:'):
        sid = int(action.rsplit(':', 1)[1])
        buttons = [[InlineKeyboardButton(c, callback_data=f"setcur:{sid}:{c}")]
                   for c in ('USD', 'BYN', 'RUB', 'EUR')]
        await query.message.reply_text("Выбери валюту:", reply_markup=InlineKeyboardMarkup(buttons))
    elif action.startswith('settings:vat:'):
        sid = int(action.rsplit(':', 1)[1])
        options = [('without_vat', 'Без НДС'), ('with_vat', 'С НДС')]
        buttons = [[InlineKeyboardButton(label, callback_data=f"setvat:{sid}:{value}")]
                   for value, label in options]
        await query.message.reply_text("Выбери статус НДС:", reply_markup=InlineKeyboardMarkup(buttons))
    elif action.startswith('setcur:') or action.startswith('setvat:'):
        kind, sid, value = action.split(':', 2)
        sid = int(sid)
        s = update_supplier_settings(sid, **({'currency': value} if kind == 'setcur' else {'tax_status': value}))
        await query.message.reply_text(f"✅ Настройки {s['name']}: {s['currency']}, {s['tax_status']}")
        await show_pending(query.message, context)
    elif action.startswith('xm:'):
        await excel_manual_callback(query, context, action)
    elif action == 'excel:cancel':
        context.user_data.pop('excel_pending', None)
        await query.message.reply_text('Excel-импорт отменён. Данные не сохранены.')
    elif action == 'excel:save':
        staged = context.user_data.get('excel_pending')
        if not staged or time.time() > staged['expires']:
            await query.message.reply_text('Подтверждение истекло. Отправь Excel ещё раз.')
            return
        sid = context.user_data.get('supplier_id')
        if sid != staged['supplier_id'] or not staged_excel_ready(staged):
            await query.message.reply_text('Поставщик изменился. Отправь Excel ещё раз.')
            return
        await query.message.reply_text('⏳ Сохраняю Excel в PostgreSQL. Подожди...')
        try:
            def _process_import():
                # Re-parse the exact staged file; no Telegram re-download required.
                extracted, skipped = [], 0
                selected_items = [it for it in staged['sheets'] if it.get('enabled')]
                for selected in selected_items:
                    batch, missing = extract_xlsx(selected.get('data', staged['data']), selected['sheet'], selected['mapping'], selected['header_row'], default_brand=_supplier_default_brand(sid))
                    extracted.extend(batch)
                    skipped += missing
                parsed = [(parse_product(name), price) for name, price in extracted]
                report = save_excel_rows(sid, parsed, staged['currency'], staged['tax_status'])
                return report, selected_items, skipped
            report, selected_items, skipped = await asyncio.to_thread(_process_import)
            context.user_data.pop('excel_pending', None)
            try:
                def _save_profiles():
                    profiles = [{**it, 'fingerprint': sheet_fingerprint(it.get('data', staged['data']), it['sheet'], it['header_row'], it['mapping'])} for it in selected_items]
                    save_excel_profiles(sid, profiles)
                await asyncio.to_thread(_save_profiles)
            except Exception:
                logger.exception('Excel profiles save failed')
        except Exception:
            logger.exception('Excel save failed')
            await query.message.reply_text('❌ Импорт не выполнен. Проверь Railway Logs. Повтори отправку файла.')
            return
        await query.message.reply_text(
            f"✅ Excel сохранён!\nФайлов: {len(staged.get('files', [1]))}\nКонфигураций: {report['saved']}\n"
            f"Нераспознанных строк (сохранены по исходному названию): {report['unrecognized_rows']}\n"
            f"Дубли внутри файла: {report['duplicates']}\n"
            f"Строк без цены: {skipped}\n"
            'Предыдущий прайс этого поставщика полностью заменён.\n'
            f"Валюта: {staged['currency']}, НДС: {staged['tax_status']}."
        )
    elif action == 'import:cancel':
        context.user_data.pop('pending', None)
        await query.message.reply_text("Прайс отменён. В базу ничего не записано.")
    elif action == 'import:save':
        pending = context.user_data.get('pending')
        sid = context.user_data.get('supplier_id')
        if not pending or not sid:
            await query.message.reply_text('Нет ожидающего прайса или поставщика.')
            return
        rows = [(parse_product(name), price) for name, price in pending]
        summary = save_price_list(sid, rows)
        context.user_data.pop('pending', None)
        await query.message.reply_text(
            f"✅ Сохранено конфигураций: {summary['saved']}\n"
            f"Дублей внутри прайса: {summary['duplicates']}\n"
            f"Пропущено для проверки: {summary['skipped']}\n"
            "Предыдущий прайс этого поставщика полностью заменён."
            + ("\nПримеры пропусков:\n" + '\n'.join(
                f"• {name[:70]} — {reason}" for name, reason in summary['rejected'])
               if summary['rejected'] else "")
        )


async def handle_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await authorize(update):
        return
    text = (update.effective_message.text or '').strip()
    delivery_sid = context.user_data.get('await_delivery_id')
    if delivery_sid is not None:
        if text.casefold() in ('отмена', '/cancel', '/cancelrate'):
            context.user_data.pop('await_delivery_id', None)
            await update.effective_message.reply_text('Настройка срока доставки отменена.')
            return
        try:
            if text.casefold() in ('нет', 'не указан', '-'):
                days = None
            elif re.fullmatch(r'\d{1,3}', text) and int(text) <= 365:
                days = int(text)
            else:
                raise ValueError()
            supplier = set_supplier_delivery_days(delivery_sid, days)
        except ValueError:
            await update.effective_message.reply_text('Введи число от 0 до 365 или «нет» для неизвестного срока.')
            return
        context.user_data.pop('await_delivery_id', None)
        await update.effective_message.reply_text(
            f"✅ {supplier['name']}: 🚚 {_delivery_label(days)}", reply_markup=settings_keyboard(supplier))
        return
    if context.user_data.get('await_fx_rate'):
        if text.lower() in ('/cancelrate', 'отмена'):
            context.user_data.pop('await_fx_rate', None)
            await update.effective_message.reply_text('Изменение курса отменено.')
            return
        try:
            value = Decimal(text.replace(',', '.').strip())
            rate = set_usd_byn_rate(value)
        except (InvalidOperation, ValueError):
            await update.effective_message.reply_text('Некорректный курс. Введи число, например 3,05.')
            return
        context.user_data.pop('await_fx_rate', None)
        await update.effective_message.reply_text(f'✅ Новый курс: 1 USD = {rate} BYN')
        return
    rename_sid = context.user_data.get('await_rename_id')
    if rename_sid is not None:
        if not text or len(text) > 120 or '\n' in text:
            await update.effective_message.reply_text('Название должно быть одной строкой (до 120 символов).')
            return
        try:
            renamed = rename_supplier(rename_sid, text)
        except ValueError as exc:
            await update.effective_message.reply_text(str(exc))
            return
        if renamed is None:
            context.user_data.pop('await_rename_id', None)
            await update.effective_message.reply_text('Поставщик больше не существует.')
            return
        context.user_data.pop('await_rename_id', None)
        await update.effective_message.reply_text(
            f"✅ Название изменено на «{renamed['name']}».",
            reply_markup=supplier_manage_keyboard(renamed)
        )
        return
    if context.user_data.pop('await_name', False):
        if not text or len(text) > 120 or '\n' in text:
            context.user_data['await_name'] = True
            await update.effective_message.reply_text('Название должно быть одной строкой (до 120 символов).')
            return
        sid = add_supplier(text)
        context.user_data['supplier_id'] = sid
        await update.effective_message.reply_text(f'✅ Поставщик «{text}» добавлен/выбран.')
        await show_pending(update.effective_message, context)
        return
    existing_draft = context.user_data.get('auto_draft')
    if existing_draft is not None:
        sid = context.user_data.get('supplier_id')
        if sid != existing_draft['supplier_id']:
            await update.effective_message.reply_text('Черновик относится к другому поставщику.', reply_markup=auto_draft_keyboard(existing_draft))
            return
        try:
            draft, added = append_to_auto_draft(context, sid, text)
        except ValueError as exc:
            await update.effective_message.reply_text(str(exc), reply_markup=auto_draft_keyboard(existing_draft))
            return
        if added:
            await auto_draft_status(update.effective_message, draft, context)
        else:
            await update.effective_message.reply_text('Такое сообщение уже есть в черновике.', reply_markup=auto_draft_keyboard(draft))
        return
    batch = context.user_data.get('parts')
    if batch is not None:
        if batch['supplier_id'] != context.user_data.get('supplier_id'):
            await update.effective_message.reply_text('Сбор относится к другому поставщику. Отмени сбор и начни заново.', reply_markup=multipart_keyboard())
            return
        if batch['messages'] >= 100 or len(batch['rows']) >= 20000:
            await update.effective_message.reply_text('Достигнут лимит частей/строк. Заверши текущий сбор.', reply_markup=multipart_keyboard())
            return
        rows, next_brand = parse_price_list(text, batch.get('brand'), return_brand=True)
        if not rows and next_brand == batch.get('brand'):
            await update.effective_message.reply_text('В этой части не удалось распознать пары «товар — цена». Часть не добавлена. Проверь формат и отправь снова.', reply_markup=multipart_keyboard())
            return
        batch['rows'].extend(rows)
        batch['brand'] = next_brand
        batch['messages'] += 1
        await update.effective_message.reply_text(
            f'✅ Часть {batch["messages"]} принята: {len(rows)} строк. '
            f'Всего накоплено: {len(batch["rows"])} строк.\n'
            'Отправляй следующую часть или заверши сбор.', reply_markup=multipart_keyboard())
        return
    current_mode = context.user_data.get('mode', 'auto')
    if current_mode == 'search':
        await run_search(update.effective_message, text, context)
        return

    rows = parse_price_list(text, _supplier_default_brand(context.user_data.get('supplier_id')))
    if current_mode == 'import' or (len(rows) >= 2 and context.user_data.get('supplier_id')):
        sid = context.user_data.get('supplier_id')
        if not sid or not get_supplier(sid):
            await update.effective_message.reply_text('Сначала выбери поставщика через /suppliers.')
            return
        try:
            draft, _ = append_to_auto_draft(context, sid, text)
        except ValueError as exc:
            await update.effective_message.reply_text(str(exc))
            return
        await auto_draft_status(update.effective_message, draft, context)
    elif len(rows) == 1 or ('\n' in text and not rows):
        context.user_data['undecided_text'] = text
        await update.effective_message.reply_text('Это поиск товара или загрузка прайса?', reply_markup=decision_keyboard())
    else:
        await run_search(update.effective_message, text, context)



def _inspect_excel_file(payload, filename, tax_status, profiles):
    """Validate one file independently; stage all its worksheets without writing offers."""
    detected = inspect_xlsx(payload, tax_status=tax_status)
    sheets = []
    for result in detected:
        name = result['sheet']
        if result['status'] == 'preview' and result.get('valid', 0) > 0 and not result.get('ambiguous') and not result.get('truncated'):
            item = dict(result, enabled=True, skipped=False)
        else:
            item = {'sheet': name, 'mapping': {k: None for k in EXCEL_FIELDS},
                    'header_row': 0, 'enabled': False, 'skipped': False}
            old = profiles.get(name)
            if old:
                oldmap, oldhead = old['mapping'], old['header_row']
                if sheet_fingerprint(payload, name, oldhead, oldmap) == old['column_fingerprint']:
                    check = inspect_mapping(payload, name, oldmap, oldhead)
                    if check['valid'] and not check['truncated']:
                        item.update(check, mapping=oldmap, header_row=oldhead, enabled=True)
        old = profiles.get(name)
        if old and not (result['status'] == 'preview' and result.get('valid', 0) > 0):
            oldmap, oldhead = old['mapping'], old['header_row']
            if sheet_fingerprint(payload, name, oldhead, oldmap) == old['column_fingerprint']:
                check = inspect_mapping(payload, name, oldmap, oldhead)
                if check['valid'] and not check['truncated']:
                    item.update(check, mapping=oldmap, header_row=oldhead, enabled=True, skipped=False)
        item.update(file_name=filename, data=payload)
        sheets.append(item)
    if not sheets:
        raise ValueError(f'В файле {filename} нет доступных листов')
    return sheets


async def _finalize_excel_batch(context, message, album_key=None):
    """One common review for a Telegram media group (or one standalone file)."""
    draft = context.user_data.get('excel_album')
    if not draft or (album_key is not None and draft['key'] != album_key):
        return
    context.user_data.pop('excel_album', None)
    sid = draft['supplier_id']
    supplier = get_supplier(sid)
    if not supplier or supplier['currency'] != draft['currency'] or supplier['tax_status'] != draft['tax_status']:
        await message.reply_text('Настройки поставщика изменились. Повтори загрузку файлов.')
        return
    if draft['errors']:
        await message.reply_text('❌ Общий Excel-прайс не принят. Старые товары сохранены.\n' + '\n'.join(draft['errors'])[:3000])
        return
    if not draft['files']:
        await message.reply_text('Не получено корректных Excel-файлов.')
        return
    sheets = [item for file in draft['files'] for item in file['sheets']]
    context.user_data['excel_pending'] = {
        'data': draft['files'][0]['data'], 'files': draft['files'],
        'sheets': sheets, 'supplier_id': sid, 'supplier_name': supplier['name'],
        'currency': supplier['currency'], 'tax_status': supplier['tax_status'],
        'expires': time.time()+1800,
    }
    await excel_review(message, context)


async def _finish_excel_album(context, message, key):
    """Telegram does not report the album length; debounce incoming updates."""
    while True:
        draft = context.user_data.get('excel_album')
        if not draft or draft['key'] != key:
            return
        version = draft['version']
        await asyncio.sleep(2.5)
        draft = context.user_data.get('excel_album')
        if not draft or draft['key'] != key:
            return
        if draft['version'] == version:
            await _finalize_excel_batch(context, message, key)
            return


_csv_import_lock = asyncio.Lock()


async def handle_large_csv(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Use a path provided by a co-located local Bot API; never buffer 350 MB."""
    if not await authorize(update):
        return
    message = update.effective_message
    document = message.document
    sid = context.user_data.get('supplier_id')
    supplier = get_supplier(sid) if sid else None
    if not supplier:
        await message.reply_text('Сначала выбери поставщика через /suppliers.')
        return
    if context.user_data.get('auto_draft') or context.user_data.get('parts') or context.user_data.get('excel_album'):
        await message.reply_text('Сначала заверши предыдущую загрузку прайса.')
        return
    if supplier['currency'] != 'BYN' or supplier['tax_status'] not in ('with_vat', 'without_vat'):
        await message.reply_text('CSV с розничной ценой в BYN: выбери для поставщика валюту BYN и НДС в /settings.')
        return
    if document.file_size and document.file_size > MAX_CSV_BYTES:
        await message.reply_text('CSV превышает установленный предел MAX_CSV_BYTES.')
        return
    if not os.getenv('TELEGRAM_LOCAL_API_URL'):
        await message.reply_text('Большой CSV требует локального Telegram Bot API (TELEGRAM_LOCAL_API_URL).')
        return
    if _csv_import_lock.locked():
        await message.reply_text('Сейчас выполняется другой CSV-импорт. Повтори позже.')
        return
    async with _csv_import_lock:
        status = await message.reply_text('📥 CSV принят. Получаю файл и импортирую в БД. Старые цены сохранятся при ошибке.')
        try:
            tg_file = await document.get_file(read_timeout=120)
            file_path = Path(tg_file.file_path or '')
            if not file_path.is_absolute() or not file_path.is_file():
                raise ValueError('Локальный Bot API не предоставил путь к файлу. Проверь совместный контейнер и --local.')
            # Local API returns a path inside its storage directory.
            root = Path(os.getenv('TELEGRAM_API_FILES_ROOT', '/data/bot-api')).resolve()
            if not file_path.resolve().is_relative_to(root):
                raise ValueError('Небезопасный путь файла от Bot API.')
            result = await asyncio.to_thread(import_large_csv, file_path, sid, supplier['tax_status'])
            await status.edit_text('✅ CSV загружен для выбранного поставщика.\n'
                f'Строк: {result["total"]:,}\nСохранено: {result["saved"]:,}\n'
                f'Пропущено: {result["skipped"]:,}\nДубликаты: {result["duplicates"]:,}\n'
                'Цена: «Розничная цена», BYN.')
        except Exception as exc:
            logger.exception('Large CSV import failed')
            await status.edit_text('❌ CSV не импортирован, старый прайс сохранён.\n' + str(exc)[:350])


async def handle_excel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await authorize(update):
        return
    message = update.effective_message
    if message.document and (message.document.file_name or '').lower().endswith('.csv'):
        await handle_large_csv(update, context)
        return
    if context.user_data.get('auto_draft'):
        await message.reply_text('Сначала сохрани или отмени текстовый черновик.', reply_markup=auto_draft_keyboard(context.user_data['auto_draft']))
        return
    if context.user_data.get('parts'):
        await message.reply_text('Сначала заверши или отмени сбор текстового прайса.', reply_markup=multipart_keyboard())
        return
    sid = context.user_data.get('supplier_id')
    supplier = get_supplier(sid) if sid else None
    if not supplier:
        await message.reply_text('Сначала выбери поставщика через /suppliers.')
        return
    if supplier['currency'] not in ('USD', 'BYN', 'RUB', 'EUR') or supplier['tax_status'] not in ('with_vat', 'without_vat'):
        await message.reply_text('Сначала укажи валюту и НДС в /settings.')
        return
    album_id = message.media_group_id
    key = f'{message.chat_id}:{album_id}' if album_id else None
    draft = context.user_data.get('excel_album')
    if draft and draft['key'] != key:
        # Never interleave separate uploads or suppliers.
        await message.reply_text('Сначала закончи текущую загрузку Excel, затем отправь следующую.')
        return
    if not draft:
        # Invalidate an older confirmation when a new upload begins.
        context.user_data.pop('excel_pending', None)
        draft = {'key': key, 'supplier_id': sid, 'currency': supplier['currency'],
                 'tax_status': supplier['tax_status'], 'files': [], 'errors': [],
                 'version': 0, 'seen': set()}
        context.user_data['excel_album'] = draft
    doc = message.document
    draft['version'] += 1
    if doc and doc.file_unique_id not in draft['seen']:
        draft['seen'].add(doc.file_unique_id)
        if not doc.file_name or not doc.file_name.lower().endswith(('.xlsx', '.xls')):
            draft['errors'].append(f'{doc.file_name or "Без имени"}: нужен файл .xlsx или .xls')
        elif doc.file_size and doc.file_size > MAX_FILE_BYTES:
            draft['errors'].append(f'{doc.file_name}: превышен лимит размера файла')
        else:
            try:
                tg_file = await doc.get_file()
                payload = bytes(await tg_file.download_as_bytearray())
                profiles = get_excel_profiles(sid)
                sheets = _inspect_excel_file(payload, doc.file_name, supplier['tax_status'], profiles)
                draft['files'].append({'name': doc.file_name, 'data': payload, 'sheets': sheets})
            except Exception as exc:
                logger.exception('Excel inspection failed: %s', doc.file_name)
                draft['errors'].append(f'{doc.file_name}: {str(exc)[:140]}')
    if album_id:
        if draft['version'] == 1:
            context.application.create_task(_finish_excel_album(context, message, key), update=update)
    else:
        await _finalize_excel_batch(context, message)


async def find(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await authorize(update):
        return
    query_text = " ".join(context.args or []).strip()
    if not query_text:
        await update.effective_message.reply_text(
            'Просто напиши название товара, например: айфон 17 про 256 чёрный'
        )
        return
    await run_search(update.effective_message, query_text, context)


def _search_normalize(text):
    text = text.casefold().replace('ё', 'е')
    text = re.sub(r'айфона?|айфон', 'iphone', text)
    text = re.sub(r'(\d+)\s*(?:гб|gb)\b', r'\1gb', text)
    return re.findall(r'[a-zа-я0-9]+', text)


def _offer_matches(query, requested, offer):
    original = offer['original_name']
    found = parse_product(original)
    if requested['product_role'] != found['product_role']:
        return False

    # For families with generations, an unspecified generation means any one.
    query_family = requested.get('family')
    if query_family:
        if found.get('family') != query_family:
            return False
        requested_gen = requested.get('generation')
        if requested_gen is not None and found.get('generation') != requested_gen:
            return False
        # Accessory/component queries also need the specified subtype.
        if requested['product_role'] != 'device':
            accessory_terms = ('левое ухо', 'правое ухо', 'левый наушник',
                               'правый наушник', 'charging case', 'зарядный кейс',
                               'зарядный футляр', 'чехол', 'case', 'cover')
            requested_kind = [x for x in accessory_terms if x in query.casefold()]
            if requested_kind and not any(x in original.casefold() for x in requested_kind):
                return False
    query_model = None if query_family else requested.get('model')
    # A stored, explicitly recognized model is valid even if the original text
    # is an abbreviation. This supports older rows that retain structured data.
    found_model = found.get('model') or offer.get('model')
    if query_model:
        if not found_model and query_model.startswith('iPhone '):
            # A bare "17 Pro" could be iPhone *or* Redmi. Only infer Apple
            # for a supplier explicitly designated as Apple-only; never infer
            # it from a matching number alone.
            if _supplier_default_brand(offer.get('supplier_id')) == 'Apple':
                inferred = contextual_name(original, 'Apple')
                found = parse_product(inferred)
                found_model = found.get('model')
        if not found_model or query_model.casefold() != found_model.casefold():
            return False
    elif not query_family:
        if not all(t in _search_normalize(original) for t in _search_normalize(query)):
            return False

    for field in ('storage_gb', 'ram_gb', 'color', 'region', 'sim'):
        required = requested.get(field)
        if required is not None and (found.get(field) or offer.get(field)) != required:
            return False
    return True


def _display(value, labels=None):
    if value is None or value == '' or value == []:
        return 'Не указан'
    return labels.get(value, value) if labels else str(value)


SIM_LABELS = {'esim': 'eSIM', 'sim_esim': 'SIM + eSIM',
              'dual_sim': '2 SIM', 'dual_esim': '2 eSIM'}
COLOR_LABELS = {'black': 'Black', 'white': 'White', 'blue': 'Blue',
                'deep_blue': 'Deep Blue', 'sky_blue': 'Sky Blue',
                'silver': 'Silver', 'green': 'Green', 'orange': 'Orange',
                'purple': 'Purple', 'cosmic_orange': 'Cosmic Orange',
                'natural_titanium': 'Natural Titanium', 'burgundy': 'Burgundy',
                'glacier': 'Glacier', 'starlight': 'Starlight',
                'midnight': 'Midnight', 'graphite': 'Graphite',
                'gold': 'Gold', 'pink': 'Pink'}


# TV/monitor diagonals are often encoded in the model, not in a separate field.
SCREEN_DIAGONALS = {
    'tv': {24, 28, 32, 40, 42, 43, 48, 49, 50, 55, 58, 60, 65, 70, 75, 77, 85, 86, 98, 100, 115},
    'monitor': {15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 27, 28, 29, 30, 31, 32, 34, 35, 38, 40, 42, 43, 45, 49, 57},
}
SCREEN_PAGE_SIZE = 10


def _screen_category(text):
    t = str(text).casefold().replace('ё', 'е')
    if re.search(r'(?<!\w)(?:телевизор\w*|телик\w*|tv)(?!\w)', t):
        return 'tv'
    if re.search(r'(?<!\w)(?:монитор\w*|monitor\w*)(?!\w)', t):
        return 'monitor'
    return None


def _screen_diagonal(name, category):
    """Recognize explicit inch sizes or conservative, manufacturer-style model prefixes."""
    t = str(name).upper().replace('С', 'C')  # Cyrillic C occurs in some SKUs.
    allowed = SCREEN_DIAGONALS[category]
    # Explicit 27-inch / 27″ / 27 дюймов; no bare arbitrary digits.
    for pattern in (r'(?<!\d)(\d{2,3})\s*(?:[\"″]|ДЮЙМ\w*|INCH(?:ES)?)(?!\w)',
                    r'(?<!\w)(?:ДИАГОНАЛЬ|DIAGONAL)\s*[:=-]?\s*(\d{2,3})(?!\d)'):
        for match in re.finditer(pattern, t):
            value = int(match.group(1))
            if value in allowed:
                return value
    # Xiaomi L75MC..., Samsung UE55... / QE65..., LG 55U..., TCL 55C..., etc.
    models = ([r'(?<![A-Z0-9])L(\d{2,3})[A-Z]',
               r'(?<![A-Z0-9])(?:UE|QE|QN|Q[E]?|KD|XR|K|OLED)(\d{2,3})[A-Z0-9]',
               r'(?<![A-Z0-9])(\d{2,3})(?:[A-Z]{1,3}\d|[A-Z]{2,})']
              if category == 'tv' else
              [r'(?<![A-Z0-9])(?:LS|27|M|G|S|P|VG|VA|XG|XV|AW|U)(\d{2})[A-Z]',
               r'(?<![A-Z0-9])(\d{2})(?:[A-Z]{2,}\d|[A-Z]{2,})'])
    for pattern in models:
        for match in re.finditer(pattern, t):
            value = int(match.group(1))
            if value in allowed:
                return value
    return None


def _screen_query_parts(text, category):
    t = re.sub(r'(?i)(?<!\w)(?:телевизор\w*|телик\w*|монитор\w*|monitor\w*|tv)(?!\w)', ' ', text)
    size = re.search(r'(?<!\w)(\d{2,3})\s*(?:[\"″]|дюйм\w*|inch(?:es)?|(?=\s|$))', t, re.I)
    diagonal = int(size.group(1)) if size and int(size.group(1)) in SCREEN_DIAGONALS[category] else None
    if diagonal is not None:
        t = t[:size.start()] + ' ' + t[size.end():]
    return diagonal, _search_normalize(t)


def _screen_offer_matches(offer, category, terms):
    name = offer['original_name']
    if _screen_category(name) != category:
        return False
    tokens = _search_normalize(name)
    return all(token in tokens for token in terms)


def _screen_diagonal_keyboard(token, matches, selected=None):
    counts = {}
    for offer in matches:
        d = _screen_diagonal(offer['original_name'], _screen_category(offer['original_name']))
        counts[d] = counts.get(d, 0) + 1
    sizes = sorted(d for d in counts if d is not None)
    rows = []
    for i in range(0, len(sizes), 3):
        rows.append([InlineKeyboardButton(f'{d}″ ({counts[d]})' + (' ✓' if selected == d else ''),
                                        callback_data=f'sr:diag:{token}:{d}') for d in sizes[i:i+3]])
    if None in counts:
        rows.append([InlineKeyboardButton(f'❔ Не определена ({counts[None]})', callback_data=f'sr:diag:{token}:unknown')])
    rows.append([InlineKeyboardButton(f'📋 Все диагонали ({len(matches)})', callback_data=f'sr:diag:{token}:any')])
    return InlineKeyboardMarkup(rows)


def format_offer(offer, index, rate=None):
    parsed = parse_product(offer['original_name'])
    def get(key):
        return parsed.get(key) or offer.get(key)
    memory = get('storage_gb')
    memory_label = f'{memory} GB' if memory is not None else 'Не указан'
    try:
        price = Decimal(str(offer['price']))
    except (InvalidOperation, TypeError, ValueError):
        price = None
    currency = (offer.get('currency') or '').upper()
    if price is not None and currency == 'BYN':
        price_label = f"{price.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)} BYN"
    elif price is not None and currency == 'USD' and rate is not None and rate > 0:
        byn = (price * rate).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        price_label = f"{price.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)} USD ({byn} BYN)"
    else:
        # No configured RUB/EUR conversion rate: never mislabel foreign prices as BYN.
        price_label = f"{offer['price']} {currency or 'валюта не указана'}"
    category = _screen_category(offer['original_name'])
    if category:
        diagonal = _screen_diagonal(offer['original_name'], category)
        specs = f"📺 Диагональ: {diagonal}″" if diagonal else "📺 Диагональ: не определена"
    else:
        specs = (f"Память: {memory_label} | Цвет: {_display(get('color'), COLOR_LABELS)}\n"
                 f"SIM: {_display(get('sim'), SIM_LABELS)} | Регион: {_display(get('region'))}")
    return (f"{index}. 🏪 {offer['supplier_name']}\n"
            f"{offer['original_name']}\n"
            f"💰 {price_label}\n"
            f"{specs}")

def _interpret_search_query(query_text):
    """Short iPhone-style queries default to iPhone, never to Redmi/Xiaomi.

    An explicit brand always wins. Other bare model families (e.g. S25) keep
    their existing behavior. Matching still requires a recognized iPhone model.
    """
    compact = re.sub(r"\s+", " ", query_text.strip())
    if re.match(r"^\d{1,2}\s+(?:pro\s*max|promax|pro|plus|mini|air)(?=\s|$)", compact, re.I):
        return "iPhone " + compact
    return compact


def _usd_equivalent(offer, rate):
    try:
        price = Decimal(str(offer['price']))
    except (InvalidOperation, TypeError, ValueError):
        return None
    currency = (offer.get('currency') or '').upper()
    if currency == 'USD':
        return price
    if currency == 'BYN' and rate > 0:
        return price / rate
    return None


def _offer_sort_key(offer, rate, mode='delivery'):
    days = offer.get('delivery_workdays')
    amount = _usd_equivalent(offer, rate)
    delivery_key = (days is None, days if days is not None else 9999)
    price_key = (amount is None, amount if amount is not None else Decimal('Infinity'))
    tie = ((offer.get('currency') or '').upper(),
           (offer.get('supplier_name') or '').casefold(),
           (offer.get('original_name') or '').casefold())
    return (*price_key, *delivery_key, *tie) if mode == 'price' else (*delivery_key, *price_key, *tie)


def _sort_choice_keyboard(token):
    return InlineKeyboardMarkup([[
        InlineKeyboardButton('🚚 По доставке', callback_data=f'sr:delivery:{token}:0'),
        InlineKeyboardButton('💰 По цене', callback_data=f'sr:price:{token}:0'),
    ]])


def _render_offer(offer, index, rate):
    result = format_offer(offer, index, rate)
    days = offer.get('delivery_workdays')
    when = _delivery_date(days)
    result += f'\n🚚 Доставка: {_delivery_label(days)}'
    if when is not None:
        result += f' (ориентировочно {when})'
    return result


def _search_nav_keyboard(token, page, pages, mode, screen=False):
    rows = [[InlineKeyboardButton('🚚 Доставка' + (' ✓' if mode == 'delivery' else ''),
                                  callback_data=f'sr:delivery:{token}:0'),
             InlineKeyboardButton('💰 Цена' + (' ✓' if mode == 'price' else ''),
                                  callback_data=f'sr:price:{token}:0')]]
    navigation = []
    if page > 0:
        navigation.append(InlineKeyboardButton('◀️ Назад', callback_data=f'sr:page:{token}:{page-1}'))
    navigation.append(InlineKeyboardButton(f'{page+1}/{pages}', callback_data=f'sr:page:{token}:{page}'))
    if page + 1 < pages:
        navigation.append(InlineKeyboardButton('Вперёд ▶️', callback_data=f'sr:page:{token}:{page+1}'))
    rows.append(navigation)
    if screen:
        rows.append([InlineKeyboardButton('📺 Выбрать диагональ', callback_data=f'sr:diags:{token}:0')])
    return InlineKeyboardMarkup(rows)


async def _send_search_blocks(message, header, offers, rate, offset=0, keyboard=None):
    # Telegram limit is 4096 characters; split long pages into safe messages.
    blocks = [header] + [_render_offer(offer, i + offset + 1, rate) for i, offer in enumerate(offers)]
    chunks, chunk = [], ''
    for block in blocks:
        if len(block) > 3600:
            block = block[:3595] + '…'
        if chunk and len(chunk) + len(block) + 2 > 3800:
            chunks.append(chunk)
            chunk = ''
        chunk += ('\n\n' if chunk else '') + block
    if chunk:
        chunks.append(chunk)
    for i, value in enumerate(chunks):
        await message.reply_text(value, reply_markup=keyboard if i == len(chunks) - 1 else None)


def _search_prefilter_anchor(text, requested):
    """Choose one safe substring to cut DB->Python traffic.

    Prefer distinctive model numbers/family names. Using one fragment (not
    all tokens) preserves abbreviations and supplier-specific descriptions.
    No good fragment => old exhaustive behavior.
    """
    model = (requested.get('model') or '')
    family = (requested.get('family') or '')
    if family:
        if 'airpods' in family.casefold():
            return 'airpods'
        if 'watch' in family.casefold():
            return 'watch'
    if model:
        # '17' matches both abbreviated '17 Pro' and full iPhone names.
        candidates = re.findall(r'(?<!\w)\d{2,3}[a-z]?(?!\w)', model.casefold())
        if candidates:
            return candidates[0]
        if 'iphone air' in model.casefold():
            return 'air'
        words = re.findall(r'[a-zа-я]{4,}', model.casefold())
        if words:
            return words[0]
    # For unknown categories use a literal term only when it is distinctive;
    # short and ambiguous queries are deliberately scanned exhaustively.
    words = re.findall(r'[a-zа-я0-9]+', text.casefold())
    words = [w for w in words if len(w) >= 4 and w not in {'pro', 'max', 'plus', 'with', 'case'}]
    return max(words, key=len) if words else None


async def run_search(message, query_text, context):
    query_text = query_text.strip()
    if not query_text:
        await message.reply_text('Напиши название товара для поиска.')
        return
    category = _screen_category(query_text)
    diagonal = None
    if category:
        diagonal, terms = _screen_query_parts(query_text, category)
        anchor = 'телевиз' if category == 'tv' else 'монитор'
        matches = [o for o in search_offers(anchor) if _screen_offer_matches(o, category, terms)]
    else:
        interpreted = _interpret_search_query(query_text)
        requested = parse_product(interpreted)
        anchor = _search_prefilter_anchor(query_text, requested)
        model = requested.get('model') or ''
        # All recognized iPhone generations use the indexed model column.
        # Unknown/abbreviated supplier names are also considered by requiring
        # both generation and storage where specified. Python validates the
        # complete model, capacity and color afterwards.
        iphone = re.fullmatch(
            r'iPhone (\d{1,2})(?: (Pro Max|Pro|Plus|Mini|Air))?', model, re.I
        )
        if iphone:
            generation = iphone.group(1)
            variant = iphone.group(2)
            storage = requested.get('storage_gb')
            if variant:
                # Pro Max previously had this optimized path; extend it to
                # other iPhone variants without requiring supplier branding.
                variant_terms = ('pro', 'max') if variant.casefold() == 'pro max' else (variant.casefold(),)
                name_terms = (generation, *variant_terms)
            elif storage is not None:
                # Generation alone matches far too many catalogue entries.
                # Storage is required by the final matcher, so using it as a
                # prefilter cannot discard a valid textual match.
                name_terms = (generation, str(storage))
            else:
                name_terms = (generation, 'iphone')
            candidates = search_offers(model=model, name_terms=name_terms)
        else:
            candidates = search_offers(anchor)
        matches = [o for o in candidates
                   if _offer_matches(query_text, requested, o)]
    if not matches:
        await message.reply_text(f'🔎 {query_text}\nПо актуальным прайсам совпадений нет.')
        return
    rate = get_usd_byn_rate()
    token = int(time.time_ns() % 1000000000000)
    snapshot = {
        'token': token, 'query': query_text, 'all': matches, 'rate': rate,
        'category': category, 'diagonal': diagonal, 'mode': 'price',
    }
    context.user_data['search_result_snapshot'] = snapshot
    if category and diagonal is None:
        await message.reply_text(
            f'🔎 {query_text}\nНайдено предложений: {len(matches)}\nВыбери диагональ:',
            reply_markup=_screen_diagonal_keyboard(token, matches),
        )
    else:
        await _show_search_page(message, snapshot, 0)


def _filtered_search_offers(snapshot):
    matches = snapshot['all']
    chosen = snapshot.get('diagonal')
    if chosen == 'any' or chosen is None:
        return matches
    actual = None if chosen == 'unknown' else chosen
    return [o for o in matches if _screen_diagonal(o['original_name'], snapshot['category']) == actual]


async def _show_search_page(message, snapshot, page):
    rate = snapshot['rate']
    ordered = sorted(_filtered_search_offers(snapshot),
                     key=lambda o: _offer_sort_key(o, rate, snapshot['mode']))
    pages = max(1, (len(ordered) + SCREEN_PAGE_SIZE - 1) // SCREEN_PAGE_SIZE)
    page = max(0, min(page, pages - 1))
    start = page * SCREEN_PAGE_SIZE
    mode_label = 'сроку доставки' if snapshot['mode'] == 'delivery' else 'цене'
    d = snapshot.get('diagonal')
    diagonal_label = (f'\nДиагональ: {d}″' if isinstance(d, int) else
                      '\nДиагональ: не определена' if d == 'unknown' else '')
    header = (f"🔎 {snapshot['query']}{diagonal_label}\n"
              f"Найдено предложений: {len(ordered)}\n"
              f"Страница {page+1}/{pages} · Сортировка: по {mode_label}\n"
              f"💱 1 USD = {rate} BYN")
    await _send_search_blocks(message, header, ordered[start:start+SCREEN_PAGE_SIZE], rate,
                              offset=start, keyboard=_search_nav_keyboard(
                                  snapshot['token'], page, pages, snapshot['mode'], bool(snapshot['category'])))


async def search_results_callback(query, context, action):
    parts = action.split(':')
    snapshot = context.user_data.get('search_result_snapshot')
    if (len(parts) != 4 or not snapshot or parts[2] != str(snapshot['token'])):
        await query.message.reply_text('Результаты этого поиска устарели. Повтори запрос.')
        return
    mode = parts[1]
    if mode == 'diags' and snapshot.get('category'):
        await query.message.reply_text(
            f"🔎 {snapshot['query']}\nВыбери диагональ:",
            reply_markup=_screen_diagonal_keyboard(snapshot['token'], snapshot['all'], snapshot['diagonal']),
        )
        return
    if mode == 'diag' and snapshot.get('category'):
        raw = parts[3]
        valid = {str(d) for d in SCREEN_DIAGONALS[snapshot['category']]}
        if raw not in valid | {'any', 'unknown'}:
            return
        snapshot['diagonal'] = int(raw) if raw.isdigit() else raw
        await _show_search_page(query.message, snapshot, 0)
        return
    if mode in ('delivery', 'price'):
        snapshot['mode'] = mode
        page = 0
    elif mode == 'page':
        try:
            page = max(0, int(parts[3]))
        except ValueError:
            return
    else:
        return
    await _show_search_page(query.message, snapshot, page)


async def cancelrate(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop('await_fx_rate', None)
    await update.effective_message.reply_text('Ввод курса отменён.')


async def on_error(update: object, context: ContextTypes.DEFAULT_TYPE):
    logger.error("Telegram error: %s", context.error, exc_info=context.error)


async def mts_update_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Owner-only manual MTS refresh; runs outside Telegram event loop."""
    if not is_owner(update):
        await update.effective_message.reply_text('🔒 Запуск парсера доступен только владельцу.')
        return
    import asyncio
    if os.getenv('MTS_MANUAL_ENABLED', '0').strip() != '1':
        await update.effective_message.reply_text('Ручной сбор выключен (MTS_MANUAL_ENABLED=0).')
        return
    state = context.application.bot_data.setdefault('mts_manual_state', {'running': False})
    if state['running']:
        await update.effective_message.reply_text('⏳ Сбор МТС уже запущен через Telegram.')
        return
    state['running'] = True
    chat_id = update.effective_chat.id
    await update.effective_message.reply_text('🔄 Запускаю сбор МТС: смартфоны, планшеты, часы, наушники, ноутбуки. Результат пришлю сюда.')
    async def perform():
        try:
            from web_worker import CATEGORIES, run
            result = await asyncio.to_thread(run, CATEGORIES, int(os.getenv('MTS_MAX_PAGES', '50')))
            descriptions = {
                0: '✅ Сбор МТС завершён; результаты сохранены в web_offers.',
                2: '⚠️ Сбор завершился с ошибками. Проверь Deploy Logs; часть категорий могла сохраниться.',
                3: '⏳ Другой сбор МТС уже выполняется. Новый не запущен.',
            }
            await context.bot.send_message(chat_id, descriptions.get(result, f'MTS: код завершения {result}'))
        except Exception:
            logger.exception('Manual MTS crawl failed')
            await context.bot.send_message(chat_id, '❌ Сбор МТС завершился ошибкой. Проверь Deploy Logs.')
        finally:
            state['running'] = False
    context.application.create_task(perform(), name='mts_manual_update')


async def combined_post_init(application):
    """Attach both store schedulers to the existing Telegram application."""
    await mts_post_init(application)
    await a1_bot_integration.post_init(application)


def main():
    token = os.environ['TELEGRAM_BOT_TOKEN']
    init_db()
    logger.info('PostgreSQL: подключение успешно, таблицы готовы')
    builder = ApplicationBuilder().token(token).post_init(combined_post_init)
    local_api = os.getenv('TELEGRAM_LOCAL_API_URL', '').rstrip('/')
    if local_api:
        builder = (builder.base_url(local_api + '/bot')
                   .base_file_url(local_api + '/file/bot').local_mode(True))
    app = builder.build()
    for command, handler in [('start', start), ('myid', myid), ('suppliers', suppliers),
                             ('addsupplier', addsupplier), ('settings', settings), ('prices', prices),
                             ('find', find), ('menu', menu), ('mode', mode)]:
        app.add_handler(CommandHandler(command, handler))
    app.add_handler(CommandHandler('cancelrate', cancelrate))
    app.add_handler(CommandHandler('mtsupdate', mts_update_command))
    a1_bot_integration.register(app)
    app.add_handler(CallbackQueryHandler(callbacks))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_excel))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    app.add_error_handler(on_error)
    logger.info('Mobi Price Finder запущен')
    app.run_polling()


if __name__ == '__main__':
    main()
