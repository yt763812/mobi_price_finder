"""Safe, preview-only XLSX inspection. No database writes."""
import io
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from collections import Counter
from openpyxl import load_workbook
from product_parser import parse_product, section_brand, contextual_name

NAME_HEADERS = ('модель', 'номенклатура', 'наименование', 'название товара', 'товар', 'product', 'model', 'item', 'description')
PRICE_HEADERS = ('цена', 'закупочная цена', 'опт', 'оптовая цена', 'price', 'wholesale', 'dealer price', 'цена $ usd', 'цена usd', 'цена $', 'цена rub', 'цена руб')
BRAND_HEADERS = ('бренд', 'брэнд', 'brand', 'производитель')
COLOR_HEADERS = ('цвет', 'color', 'colour')
REGION_HEADERS = ('происхождение', 'регион', 'страна', 'region')
CATEGORY_HEADERS = ('категория', 'category')
BAD_PRICE = ('итого', 'total', 'сумма', 'розница', 'retail', 'ррц')
MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_ROWS = 30000


def _clean(value):
    return str(value).strip() if value is not None else ''


def _header(value):
    return re.sub(r'\s+', ' ', _clean(value).casefold())


def _preferred_price(row):
    # Explicit USD and generic price columns take priority over RUB totals.
    indices = _indices(row, PRICE_HEADERS)
    return sorted(indices, key=lambda i: (
        0 if ('usd' in _header(row[i]) or '$' in _header(row[i])) else
        2 if ('rub' in _header(row[i]) or 'руб' in _header(row[i])) else 1, i))



def _indices(row, aliases):
    return [i for i, v in enumerate(row) if _header(v) in aliases]


def _price(value):
    """Read numeric Excel prices without mistaking float precision for invalid prices."""
    if isinstance(value, bool) or value is None:
        return None
    raw = _clean(value).replace('\u00a0', '').replace(' ', '').replace(',', '.')
    if not re.fullmatch(r'\d+(?:\.\d+)?', raw):
        return None
    try:
        result = Decimal(raw)
        if not result.is_finite() or result <= 0:
            return None
        return result.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None


def _price_column_is_suspicious(values):
    """A date copied down a column must not be accepted as a price column."""
    if len(values) < 8:
        return False
    normalized = [str(x) for x in values]
    common = Counter(normalized).most_common(1)[0][1]
    return common / len(normalized) >= 0.95 and len(set(normalized)) <= 2


def _column_price_quality(ws, mapping, header_row, max_sample=150):
    prices = []
    names = 0
    for row in ws.iter_rows(min_row=header_row + 1, max_row=min(ws.max_row or max_sample, header_row + max_sample), values_only=True):
        n = mapping.get('name'); p = mapping.get('price')
        name = _clean(row[n]) if n is not None and n < len(row) else ''
        if len(name) < 5 or not re.search(r'[a-zа-я]', name, re.I):
            continue
        names += 1
        val = _price(row[p]) if p is not None and p < len(row) else None
        if val is not None:
            prices.append(val)
    return names, prices


def _positional_mapping(ws):
    """Detect price columns by validity AND variation; repeated dates are rejected."""
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 120, 120), values_only=True))
    if not rows:
        return None
    scores = []
    for name_idx in range(min(ws.max_column or 0, 12)):
        for price_idx in range(min(ws.max_column or 0, 12)):
            if name_idx == price_idx:
                continue
            prices = []
            for row in rows:
                name = _clean(row[name_idx]) if name_idx < len(row) else ''
                val = row[price_idx] if price_idx < len(row) else None
                if len(name) >= 9 and re.search(r'[A-Za-zА-Яа-я]', name):
                    price = _price(val)
                    if price is not None:
                        prices.append(price)
            if len(prices) < 4 or _price_column_is_suspicious(prices):
                continue
            diversity = len(set(prices)) / len(prices)
            scores.append((len(prices), diversity, name_idx, price_idx))
    if not scores:
        return None
    scores.sort(key=lambda x: (x[0], x[1]), reverse=True)
    good, diversity, name, price = scores[0]
    if len(scores) > 1:
        second = scores[1]
        if second[0] >= good * 0.9 and abs(second[1] - diversity) < 0.08:
            return None  # Equally plausible columns require manual confirmation.
    return {'name': name, 'price': price, 'brand': None,
            'category': None, 'color': None, 'region': None}


def inspect_xlsx(data):
    if len(data) > MAX_FILE_BYTES:
        raise ValueError('Файл больше 12 МБ. Пришли файл поменьше.')
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            header_row = None
            mapping = None
            ambiguous = False
            for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=25, values_only=True), 1):
                names = _indices(row, NAME_HEADERS)
                prices = _preferred_price(row)
                brands = _indices(row, BRAND_HEADERS)
                categories = _indices(row, CATEGORY_HEADERS)
                colors = _indices(row, COLOR_HEADERS)
                regions = _indices(row, REGION_HEADERS)
                if names and prices:
                    header_row = row_num
                    mapping = {'name': names[0], 'price': prices[0],
                               'brand': brands[0] if brands else None,
                               'category': categories[0] if categories else None,
                               'color': colors[0] if colors else None,
                               'region': regions[0] if regions else None}
                    price_header = _clean(row[prices[0]])
                    ambiguous = len(names) > 1 or (len(prices) > 1 and _header(row[prices[0]]) == _header(row[prices[1]]))
                    break
            if mapping is None:
                mapping = _positional_mapping(ws)
                if mapping is not None:
                    header_row = 0
                    price_header = 'Без заголовка (определено по содержимому)'
            if mapping is None:
                sheets.append({'sheet': ws.title, 'status': 'unknown'})
                continue
            _, price_probe = _column_price_quality(ws, mapping, header_row)
            if _price_column_is_suspicious(price_probe):
                sheets.append({'sheet': ws.title, 'status': 'unknown'})
                continue
            count = valid = no_price = 0
            samples = []
            parsed_samples = []
            uncertain = 0
            for row in ws.iter_rows(min_row=header_row + 1, max_row=min(ws.max_row or MAX_ROWS, MAX_ROWS), values_only=True):
                def val(idx):
                    return row[idx] if idx is not None and idx < len(row) else None
                name = _clean(val(mapping['name']))
                if not name or len(name) < 5:
                    continue
                count += 1
                price = _price(val(mapping['price']))
                if price is None:
                    no_price += 1
                    continue
                valid += 1
                brand = _clean(val(mapping['brand']))
                full_name = (brand + " " + name).strip() if brand else name
                for field in ("color", "region"):
                    extra = _clean(val(mapping.get(field)))
                    if extra and extra.casefold() not in full_name.casefold():
                        full_name += " " + extra
                if len(samples) < 4:
                    samples.append(f'{full_name[:85]} — {price}')
                    parsed = parse_product(full_name)
                    parsed_samples.append({
                        'model': parsed.get('model') or 'По исходному названию',
                        'storage': parsed.get('storage_gb'),
                        'color': parsed.get('color'),
                        'region': parsed.get('region'),
                        'sim': parsed.get('sim'),
                        'warnings': parsed.get('warnings', []),
                    })
                if parse_product(full_name).get('warnings'):
                    uncertain += 1
            sheets.append({'sheet': ws.title, 'status': 'preview', 'header_row': header_row,
                           'mapping': mapping, 'ambiguous': ambiguous,
                           'price_header': price_header,
                           'named': count, 'valid': valid, 'no_price': no_price,
                           'samples': samples, 'parsed_samples': parsed_samples,
                           'uncertain': uncertain,
                           'truncated': bool(ws.max_row and ws.max_row > MAX_ROWS)})
        return sheets
    finally:
        wb.close()


def extract_xlsx(data, sheet_name, mapping, header_row, default_brand=None):
    """Extract all validated rows from one explicitly selected sheet."""
    if len(data) > MAX_FILE_BYTES:
        raise ValueError('Файл слишком большой')
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        if ws.max_row and ws.max_row > MAX_ROWS:
            raise ValueError('Превышен лимит 30000 строк')
        _, price_probe = _column_price_quality(ws, mapping, header_row)
        if _price_column_is_suspicious(price_probe):
            raise ValueError('Столбец цены содержит почти одинаковые значения (возможно, дату прайса). Выбери правильный столбец.')
        rows = []
        skipped = 0
        active_brand = default_brand
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            def val(key):
                i = mapping.get(key)
                return row[i] if i is not None and i < len(row) else None
            name = _clean(val('name'))
            heading = section_brand(name)
            if heading and _price(val('price')) is None:
                active_brand = heading
                continue
            if not name or len(name) < 5:
                continue
            amount = _price(val('price'))
            if amount is None:
                skipped += 1
                continue
            brand = _clean(val('brand'))
            full_name = (brand + ' ' + name).strip() if brand and not name.casefold().startswith(brand.casefold() + ' ') else contextual_name(name, active_brand)
            # Keep extra attributes in the searchable name: database schema remains unchanged.
            for key in ('color', 'region'):
                extra = _clean(val(key))
                if extra and extra.casefold() not in full_name.casefold():
                    full_name += ' ' + extra
            rows.append((full_name, amount))
        return rows, skipped
    finally:
        wb.close()


def sheet_columns(data, sheet_name, max_columns=40):
    """Candidate columns and small excerpts for interactive manual selection."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        width = min(ws.max_column or 0, max_columns)
        head = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 12), values_only=True))
        columns = []
        for idx in range(width):
            snippets = []
            for row in head:
                val = _clean(row[idx] if idx < len(row) else None)
                if val and val not in snippets:
                    snippets.append(val[:35])
                if len(snippets) == 2:
                    break
            columns.append({'index': idx, 'letter': __import__('openpyxl').utils.get_column_letter(idx+1), 'examples': snippets})
        return columns
    finally:
        wb.close()


def inspect_mapping(data, sheet_name, mapping, header_row):
    """Score a proposed mapping using actual priced product rows."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        if ws.max_row and ws.max_row > MAX_ROWS:
            return {'valid': 0, 'named': 0, 'no_price': 0, 'samples': [], 'truncated': True}
        if mapping.get('name') is None or mapping.get('price') is None or mapping['name'] == mapping['price']:
            return {'valid': 0, 'named': 0, 'no_price': 0, 'samples': [], 'truncated': False}
        _, price_probe = _column_price_quality(ws, mapping, header_row)
        if _price_column_is_suspicious(price_probe):
            return {'valid': 0, 'named': 0, 'no_price': 0, 'samples': [], 'truncated': False, 'warning': 'Подозрительный столбец цены: одинаковые значения, возможно дата'}
        valid = named = no_price = 0
        samples = []
        for row in ws.iter_rows(min_row=header_row+1, values_only=True):
            def v(key):
                i = mapping.get(key)
                return row[i] if i is not None and i < len(row) else None
            name = _clean(v('name'))
            if len(name) < 5 or not re.search('[a-zа-я]', name, re.I):
                continue
            named += 1
            price = _price(v('price'))
            if price is None:
                no_price += 1
                continue
            valid += 1
            if len(samples) < 3:
                samples.append(f'{name[:65]} — {price}')
        return {'valid': valid, 'named': named, 'no_price': no_price, 'samples': samples, 'truncated': False}
    finally:
        wb.close()


def sheet_fingerprint(data, sheet_name, header_row, mapping):
    """A light column signature; do not reuse a profile on a shifted/changed sheet."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        if header_row:
            row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
            return {'width': ws.max_column, 'headers': [_header(x) for x in row[:40]]}
        # Headerless profile: column alignment cannot be verified by headers;
        # later inspect_mapping validates that name and price still hold real data.
        return {'width': ws.max_column, 'headers': []}
    finally:
        wb.close()
